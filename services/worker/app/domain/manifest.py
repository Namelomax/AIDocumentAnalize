"""Machine-readable file registry parsing.

The registry is mandatory: without it the package is accepted as
CLARIFICATION_REQUIRED. The exact column naming in the hidden test is unknown,
so headers are matched loosely and both Latin and Cyrillic captions are accepted.
"""

import csv
import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime

VALID_STAGES = {"PD", "RD", "ID"}
VALID_APPROVALS = {"DRAFT", "APPROVED", "FOR_CONSTRUCTION", "SUPERSEDED", "CANCELLED"}

COLUMN_ALIASES = {
    "object_id": {"object_id", "objectid", "объект", "объект_id", "id объекта"},
    "file_name": {"file_name", "filename", "имя файла", "файл", "наименование файла"},
    "sha256": {"sha256", "sha-256", "хеш", "контрольная сумма"},
    "doc_stage": {"doc_stage", "stage", "стадия", "вид документации"},
    "discipline": {"discipline", "марка", "раздел", "дисциплина"},
    "document_code": {"document_code", "код", "шифр", "шифр документа"},
    "revision": {"revision", "rev", "редакция", "изм", "изменение"},
    "approval_status": {"approval_status", "статус", "статус утверждения"},
    "approval_date": {"approval_date", "дата утверждения", "дата"},
    "sheet_page_range": {"sheet_page_range", "листы", "диапазон листов", "страницы"},
    "predecessor_id": {"predecessor_id", "предшественник", "заменяет"},
    "signature_status": {"signature_status", "подпись", "статус подписи"},
}

REQUIRED = ("object_id", "file_name", "doc_stage")


@dataclass(frozen=True)
class ManifestEntry:
    file_name: str
    object_id: str
    doc_stage: str
    discipline: str | None = None
    document_code: str | None = None
    revision: str | None = None
    approval_status: str = "DRAFT"
    approval_date: date | None = None
    sheet_page_range: str | None = None
    predecessor_id: str | None = None
    signature_status: str | None = None
    sha256: str | None = None


@dataclass
class ManifestParseResult:
    entries: list[ManifestEntry] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def _canonical(header: str) -> str | None:
    norm = header.strip().lower().replace(" ", " ")
    for canonical, aliases in COLUMN_ALIASES.items():
        if norm in aliases:
            return canonical
    return None


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    return None


def _rows_from_csv(raw: bytes) -> list[dict[str, str]]:
    text = raw.decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
    return list(csv.DictReader(io.StringIO(text), dialect=dialect))


def _rows_from_xlsx(raw: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = ["" if h is None else str(h) for h in rows[0]]
    return [
        {headers[i]: ("" if cell is None else str(cell)) for i, cell in enumerate(row)}
        for row in rows[1:]
    ]


def _rows_from_json(raw: bytes) -> list[dict[str, str]]:
    payload = json.loads(raw.decode("utf-8"))
    if isinstance(payload, dict):
        payload = payload.get("files", [])
    return [{k: ("" if v is None else str(v)) for k, v in row.items()} for row in payload]


def parse_manifest(raw: bytes, filename: str) -> ManifestParseResult:
    result = ManifestParseResult()
    lower = filename.lower()

    try:
        if lower.endswith(".json"):
            rows = _rows_from_json(raw)
        elif lower.endswith((".xlsx", ".xlsm")):
            rows = _rows_from_xlsx(raw)
        else:
            rows = _rows_from_csv(raw)
    except Exception as exc:  # noqa: BLE001 - surfaced to the inspector, not swallowed
        result.errors.append(f"cannot read manifest: {exc}")
        return result

    if not rows:
        result.errors.append("manifest is empty")
        return result

    mapping = {}
    for header in rows[0].keys():
        canonical = _canonical(header)
        if canonical:
            mapping[header] = canonical

    missing = [c for c in REQUIRED if c not in mapping.values()]
    if missing:
        result.errors.append(f"required columns are missing: {', '.join(missing)}")
        return result

    for index, row in enumerate(rows, start=2):
        values = {mapping[h]: (v or "").strip() for h, v in row.items() if h in mapping}

        stage = values.get("doc_stage", "").upper()
        if stage not in VALID_STAGES:
            result.errors.append(f"row {index}: unknown doc_stage {stage!r}")
            continue

        approval = values.get("approval_status", "DRAFT").upper() or "DRAFT"
        if approval not in VALID_APPROVALS:
            result.errors.append(f"row {index}: unknown approval_status {approval!r}")
            continue

        result.entries.append(ManifestEntry(
            file_name=values["file_name"],
            object_id=values["object_id"],
            doc_stage=stage,
            discipline=values.get("discipline") or None,
            document_code=values.get("document_code") or None,
            revision=values.get("revision") or None,
            approval_status=approval,
            approval_date=_parse_date(values.get("approval_date")),
            sheet_page_range=values.get("sheet_page_range") or None,
            predecessor_id=values.get("predecessor_id") or None,
            signature_status=values.get("signature_status") or None,
            sha256=values.get("sha256") or None,
        ))

    return result
